"""Run VSpider agent capability regression cases.

The cases in this folder are public-site probes for generic abilities such as
forms, hover menus, tooltips, cascaders, date pickers, extraction, and paging.
They are not business dependencies: production runs do not import this module.
"""

from __future__ import annotations

import argparse
import asyncio
from dataclasses import dataclass
from dataclasses import field
from datetime import datetime
import json
from pathlib import Path
import sys
import traceback
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - exercised only when dependency missing
    load_workbook = None


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CASES = Path(__file__).with_name("cases.json")
ARTIFACTS_DIR = ROOT / "workspace" / "artifacts"
DOWNLOADS_DIR = ARTIFACTS_DIR / "downloads"
LOGS_DIR = ROOT / "logs"
REPORTS_DIR = ROOT / "workspace" / "agent_case_reports"

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from visual_web_agent.extraction_engine.snapshots import (  # noqa: E402
    DEFAULT_SNAPSHOT_DIR,
    compare_latest_snapshots,
)
from visual_web_agent.extraction_engine.strategies import (  # noqa: E402
    summarize_universal_strategies,
)
from visual_web_agent.main import run_agent  # noqa: E402
from visual_web_agent.skills.lint import lint_skill_registry  # noqa: E402
from visual_web_agent.skills.replay import check_skill_replays  # noqa: E402


@dataclass
class CaseResult:
    case_id: str
    ok: bool
    returned: bool
    log_file: str = ""
    event_stream_file: str = ""
    output_file: str = ""
    output_rows: int = 0
    error: str = ""
    diagnostics: dict[str, Any] = field(default_factory=dict)


def _load_cases(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8") as fh:
        data = json.load(fh)
    if not isinstance(data, list):
        raise ValueError(f"{path} must contain a JSON list")
    return data


def _filter_cases(
    cases: list[dict[str, Any]],
    ids: set[str] | None,
    categories: set[str] | None,
    capabilities: set[str] | None,
    *,
    include_disabled: bool = False,
) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    for case in cases:
        case_id = str(case.get("id") or "")
        case_category = str(case.get("category") or "")
        case_caps = {str(v) for v in case.get("capability", [])}
        if case.get("enabled") is False and not include_disabled and not (ids and case_id in ids):
            continue
        if ids and case_id not in ids:
            continue
        if categories and case_category not in categories:
            continue
        if capabilities and not capabilities.intersection(case_caps):
            continue
        selected.append(case)
    return selected


def _newest_created(before: set[str], folder: Path, pattern: str) -> Path | None:
    created = [
        path
        for path in folder.glob(pattern)
        if path.name not in before
    ]
    if not created:
        return None
    return max(created, key=lambda path: path.stat().st_mtime)


def _xlsx_data_rows(path: Path | None) -> int:
    if not path or not path.exists() or load_workbook is None:
        return 0
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        return max(0, int(sheet.max_row or 0) - 1)
    except Exception:
        return 0


def _read_event_stream(path: Path | None) -> list[dict[str, Any]]:
    if not path or not path.exists():
        return []
    events: list[dict[str, Any]] = []
    try:
        with path.open("r", encoding="utf-8") as fh:
            for line in fh:
                line = line.strip()
                if not line:
                    continue
                try:
                    item = json.loads(line)
                except json.JSONDecodeError:
                    events.append({"type": "parse_error", "raw": line[:300]})
                    continue
                if isinstance(item, dict):
                    events.append(item)
    except Exception as exc:
        return [{"type": "read_error", "error": repr(exc)}]
    return events


def _truncate_cell(value: Any, limit: int = 160) -> str:
    text = "" if value is None else str(value)
    text = " ".join(text.split())
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def _xlsx_summary(path: Path | None, *, preview_rows: int = 3) -> dict[str, Any]:
    summary: dict[str, Any] = {
        "path": str(path or ""),
        "rows": 0,
        "fields": [],
        "preview": [],
    }
    if not path or not path.exists() or load_workbook is None:
        return summary
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        fields = [_truncate_cell(cell, limit=80) for cell in (headers or []) if cell is not None]
        summary["fields"] = fields
        summary["rows"] = max(0, int(sheet.max_row or 0) - 1)
        preview: list[dict[str, str]] = []
        for raw_row in rows:
            row: dict[str, str] = {}
            for field_name, cell in zip(fields, raw_row):
                if field_name:
                    row[field_name] = _truncate_cell(cell)
            if any(row.values()):
                preview.append(row)
            if len(preview) >= preview_rows:
                break
        summary["preview"] = preview
    except Exception as exc:
        summary["error"] = repr(exc)
    return summary


def _event_tool_names(events: list[dict[str, Any]], key: str) -> list[str]:
    names: list[str] = []
    for event in events:
        if event.get("type") != "tool_catalog":
            continue
        for tool in event.get(key) or []:
            name = str((tool or {}).get("name") or "").strip()
            if name and name not in names:
                names.append(name)
    return names


def _extract_action_failures(events: list[dict[str, Any]]) -> list[dict[str, Any]]:
    failures: list[dict[str, Any]] = []
    for event in events:
        if event.get("type") != "act":
            continue
        result = event.get("result") or {}
        if result.get("success") is not False:
            continue
        failures.append(
            {
                "step": event.get("step"),
                "action": result.get("action", ""),
                "error": _truncate_cell(result.get("error", ""), limit=240),
                "message": _truncate_cell(result.get("message", ""), limit=240),
            }
        )
    return failures[-5:]


def _summarize_skill_diagnostics(
    *,
    skill_matches: list[dict[str, Any]],
    extracts: list[dict[str, Any]],
    guards: list[dict[str, Any]],
    verifications: list[dict[str, Any]],
) -> dict[str, Any]:
    skills = [str(item.get("skill") or "") for item in skill_matches if item.get("skill")]
    skill_sources = {str(item.get("source") or "") for item in skill_matches if item.get("source")}
    extracted_sources = {str(item.get("source") or "") for item in extracts if item.get("source")}
    verification_failures = [
        item for item in verifications
        if item.get("success") is False
    ]
    fallback_guards = [
        item for item in guards
        if "FALLBACK" in str(item.get("name") or "")
    ]
    return {
        "matched": bool(skill_matches),
        "skills": skills,
        "skill_sources": sorted(skill_sources),
        "extracted_by_skill": bool(skill_sources.intersection(extracted_sources)),
        "fallback_count": len(fallback_guards),
        "verification_failure_count": len(verification_failures),
    }


def _summarize_batch_post_extract_gate(
    guards: list[dict[str, Any]],
) -> dict[str, Any]:
    events = [
        item for item in guards
        if str(item.get("name") or "") == "BATCH_POST_EXTRACT_GATE"
    ]
    last = events[-1] if events else {}
    metadata = last.get("metadata") if isinstance(last, dict) else {}
    if not isinstance(metadata, dict):
        metadata = {}
    return {
        "triggered": bool(events),
        "count": len(events),
        "last_step": last.get("step") if last else None,
        "last_next_action": metadata.get("next_action", ""),
        "last_rows": metadata.get("rows"),
        "last_row_target": metadata.get("row_target"),
        "last_pages": metadata.get("pages"),
        "last_page_target": metadata.get("page_target"),
    }


def _last_browser_state(events: list[dict[str, Any]]) -> dict[str, Any]:
    for event in reversed(events):
        if event.get("type") != "observe":
            continue
        state = event.get("browser_state") or {}
        if not isinstance(state, dict):
            state = {}
        return {
            "step": event.get("step"),
            "url": state.get("url") or event.get("url") or "",
            "title": state.get("title", ""),
            "screenshot_path": state.get("screenshot_path") or event.get("screenshot_path") or "",
            "interactive_count": state.get("interactive_count", 0),
        }
    return {}


def build_case_diagnostics(
    *,
    event_stream_path: Path | None,
    output_path: Path | None,
) -> dict[str, Any]:
    events = _read_event_stream(event_stream_path)
    extracts = [
        {
            "step": event.get("step"),
            "source": event.get("source", ""),
            "rows": event.get("rows", 0),
            "output_file": event.get("output_file", ""),
            "metadata": event.get("metadata") or {},
        }
        for event in events
        if event.get("type") == "extract"
    ]
    guards = [
        {
            "step": event.get("step"),
            "name": event.get("name", ""),
            "message": _truncate_cell(event.get("message", ""), limit=240),
            "metadata": event.get("metadata") or {},
        }
        for event in events
        if event.get("type") == "guard"
    ]
    verifications = [
        {
            "step": event.get("step"),
            "name": event.get("name", ""),
            "success": event.get("success"),
            "observed": _truncate_cell(event.get("observed", ""), limit=200),
        }
        for event in events
        if event.get("type") == "verify"
    ]
    done_events = [
        {
            "step": event.get("step"),
            "success": event.get("success"),
            "reason": event.get("reason", ""),
            "message": _truncate_cell(event.get("message", ""), limit=240),
            "metadata": event.get("metadata") or {},
        }
        for event in events
        if event.get("type") == "done"
    ]
    dispatches = [
        {
            "step": event.get("step"),
            "action": event.get("action", ""),
            "tool": ((event.get("tool") or {}).get("name") or ""),
        }
        for event in events
        if event.get("type") == "tool_dispatch"
    ]
    skill_matches = [
        {
            "step": event.get("step"),
            "skill": ((event.get("skill") or {}).get("name") or ""),
            "action": ((event.get("skill") or {}).get("action") or ""),
            "source": ((event.get("skill") or {}).get("source") or ""),
            "capability": ((event.get("skill") or {}).get("capability") or ""),
            "score": ((event.get("metadata") or {}).get("selection") or {}).get("score"),
            "reasons": ((event.get("metadata") or {}).get("selection") or {}).get("reasons") or [],
            "history_success_rate": ((event.get("metadata") or {}).get("selection") or {}).get("history_success_rate"),
            "strategy_context": ((event.get("metadata") or {}).get("selection") or {}).get("strategy_context") or {},
        }
        for event in events
        if event.get("type") == "skill_match"
    ]
    skill_replays = [
        {
            "step": event.get("step"),
            "skill": ((event.get("skill") or {}).get("name") or ""),
            "source": ((event.get("skill") or {}).get("source") or ""),
            "path": event.get("path", ""),
            "score": ((event.get("metadata") or {}).get("score")),
        }
        for event in events
        if event.get("type") == "skill_replay"
    ]
    strategy_contexts = [
        event.get("context") or {}
        for event in events
        if event.get("type") == "strategy_context"
    ]
    run_end = next(
        (event for event in reversed(events) if event.get("type") == "run_end"),
        {},
    )
    macro_sources = [
        str(item.get("source") or "")
        for item in extracts
        if str(item.get("source") or "").endswith("_MACRO")
    ]
    skill_summary = _summarize_skill_diagnostics(
        skill_matches=skill_matches,
        extracts=extracts,
        guards=guards,
        verifications=verifications,
    )
    batch_post_extract_gate = _summarize_batch_post_extract_gate(guards)
    universal_strategies = summarize_universal_strategies(
        extracts=extracts,
        guards=guards,
    )
    return {
        "event_stream_file": event_stream_path.name if event_stream_path else "",
        "event_counts": {
            event_type: sum(1 for event in events if event.get("type") == event_type)
            for event_type in sorted({str(event.get("type") or "") for event in events})
            if event_type
        },
        "selected_tools": _event_tool_names(events, "selected_tools"),
        "skill_matches": skill_matches,
        "skill_replays": skill_replays,
        "strategy_contexts": strategy_contexts,
        "last_strategy_context": strategy_contexts[-1] if strategy_contexts else {},
        "skill_summary": skill_summary,
        "universal_strategies": universal_strategies,
        "batch_post_extract_gate": batch_post_extract_gate,
        "dispatched_tools": dispatches,
        "macro_sources": macro_sources,
        "extracts": extracts,
        "done_events": done_events[-5:],
        "last_done": done_events[-1] if done_events else {},
        "guards": guards,
        "verifications": verifications[-10:],
        "action_failures": _extract_action_failures(events),
        "last_browser_state": _last_browser_state(events),
        "run_end": {
            "success": run_end.get("success"),
            "reason": run_end.get("reason", ""),
            "metadata": run_end.get("metadata") or {},
        } if run_end else {},
        "output": _xlsx_summary(output_path),
        "download": {"path": "", "name": "", "exists": False},
    }


def _classify_failure(
    *,
    returned: bool,
    output_rows: int,
    error: str,
    expected: dict[str, Any],
    diagnostics: dict[str, Any],
) -> str:
    if error:
        if "PermissionError" in error or "拒绝访问" in error:
            return "browser_start_failed"
        return "exception"
    if any(item.get("success") is False for item in diagnostics.get("verifications", [])):
        return "verification_failed"
    expected_success = bool(expected.get("success", True))
    if returned != expected_success:
        run_end = diagnostics.get("run_end") or {}
        reason = str(run_end.get("reason") or "")
        if "max steps" in reason.lower():
            return "max_steps_reached"
        return "agent_return_mismatch"
    min_rows = expected.get("min_output_rows")
    if min_rows is not None and output_rows < int(min_rows):
        return "min_rows_not_met"
    expected_download = str(expected.get("download_file") or "").strip()
    if expected_download:
        download = diagnostics.get("download") or {}
        if download.get("name") != expected_download:
            return "download_missing"
    last_done_contains = str(expected.get("last_done_contains") or "").strip()
    if last_done_contains:
        last_done_text = json.dumps(diagnostics.get("last_done") or {}, ensure_ascii=False)
        if last_done_contains not in last_done_text:
            return "done_text_missing"
    required_fields = [str(v) for v in expected.get("required_fields") or []]
    if required_fields:
        fields = {str(v).lower() for v in ((diagnostics.get("output") or {}).get("fields") or [])}
        missing = [field for field in required_fields if field.lower() not in fields]
        if missing:
            return "field_missing"
    return ""


def _print_case(case: dict[str, Any]) -> None:
    caps = ", ".join(case.get("capability") or [])
    expected = case.get("expected") or {}
    min_rows = expected.get("min_output_rows")
    suffix = f" | min_rows={min_rows}" if min_rows else ""
    if case.get("enabled") is False:
        suffix += " | disabled"
    print(
        f"{case['id']:<32} {case.get('category',''):<12} "
        f"[{caps}]{suffix}"
    )


def _summarize_skill_usage(results: list[CaseResult]) -> dict[str, Any]:
    skill_counts: dict[str, int] = {}
    source_counts: dict[str, int] = {}
    cases_with_skill = 0
    cases_extracted_by_skill = 0
    skill_case_failures = 0
    fallback_count = 0
    verification_failure_count = 0
    for result in results:
        diagnostics = result.diagnostics or {}
        summary = diagnostics.get("skill_summary") or {}
        if not summary.get("matched"):
            continue
        cases_with_skill += 1
        if result.ok:
            pass
        else:
            skill_case_failures += 1
        if summary.get("extracted_by_skill"):
            cases_extracted_by_skill += 1
        fallback_count += int(summary.get("fallback_count") or 0)
        verification_failure_count += int(summary.get("verification_failure_count") or 0)
        for skill_name in summary.get("skills") or []:
            skill_counts[str(skill_name)] = skill_counts.get(str(skill_name), 0) + 1
        for source_name in summary.get("skill_sources") or []:
            source_counts[str(source_name)] = source_counts.get(str(source_name), 0) + 1
    success_count = max(0, cases_with_skill - skill_case_failures)
    return {
        "cases_with_skill": cases_with_skill,
        "cases_extracted_by_skill": cases_extracted_by_skill,
        "skill_case_failures": skill_case_failures,
        "skill_success_rate": (
            round(success_count / cases_with_skill, 4)
            if cases_with_skill else None
        ),
        "fallback_count": fallback_count,
        "verification_failure_count": verification_failure_count,
        "skills_used": dict(sorted(skill_counts.items())),
        "sources_used": dict(sorted(source_counts.items())),
    }


async def _run_case(case: dict[str, Any]) -> CaseResult:
    case_id = str(case["id"])
    before_logs = {path.name for path in LOGS_DIR.glob("run_log_*.html")}
    before_events = {path.name for path in LOGS_DIR.glob("event_stream_*.jsonl")}
    before_outputs = {path.name for path in ARTIFACTS_DIR.glob("*.xlsx")}
    before_downloads = {path.name for path in DOWNLOADS_DIR.glob("*")} if DOWNLOADS_DIR.exists() else set()
    returned = False
    error = ""
    try:
        returned = bool(
            await run_agent(
                str(case["url"]),
                str(case["goal"]),
                enable_xhr=bool(case.get("enable_xhr", False)),
                upload_file=str(case.get("upload_file") or ""),
            )
        )
    except Exception as exc:  # pragma: no cover - integration fallback
        error = repr(exc)
        traceback.print_exc()

    log_path = _newest_created(before_logs, LOGS_DIR, "run_log_*.html")
    event_path = _newest_created(before_events, LOGS_DIR, "event_stream_*.jsonl")
    output_path = _newest_created(before_outputs, ARTIFACTS_DIR, "*.xlsx")
    download_path = _newest_created(before_downloads, DOWNLOADS_DIR, "*") if DOWNLOADS_DIR.exists() else None
    output_rows = _xlsx_data_rows(output_path)
    diagnostics = build_case_diagnostics(
        event_stream_path=event_path,
        output_path=output_path,
    )
    diagnostics["download"] = {
        "path": str(download_path or ""),
        "name": download_path.name if download_path else "",
        "exists": bool(download_path and download_path.exists()),
    }

    expected = case.get("expected") or {}
    expected_download = str(expected.get("download_file") or "").strip()
    if expected_download and not download_path:
        fallback_download = DOWNLOADS_DIR / expected_download
        if fallback_download.exists():
            download_path = fallback_download
            diagnostics["download"] = {
                "path": str(download_path),
                "name": download_path.name,
                "exists": True,
            }
    ok = returned == bool(expected.get("success", True))
    min_rows = expected.get("min_output_rows")
    if min_rows is not None:
        ok = ok and output_rows >= int(min_rows)
    verification_failures = [
        item for item in diagnostics.get("verifications", [])
        if item.get("success") is False
    ]
    if verification_failures:
        ok = False
    if expected_download:
        ok = ok and bool(download_path and download_path.name == expected_download)
    last_done_contains = str(expected.get("last_done_contains") or "").strip()
    if last_done_contains:
        last_done_text = json.dumps(diagnostics.get("last_done") or {}, ensure_ascii=False)
        ok = ok and last_done_contains in last_done_text
    failure_type = "" if ok else _classify_failure(
        returned=returned,
        output_rows=output_rows,
        error=error,
        expected=expected,
        diagnostics=diagnostics,
    )
    diagnostics["failure_type"] = failure_type

    return CaseResult(
        case_id=case_id,
        ok=ok,
        returned=returned,
        log_file=log_path.name if log_path else "",
        event_stream_file=event_path.name if event_path else "",
        output_file=output_path.name if output_path else "",
        output_rows=output_rows,
        error=error,
        diagnostics=diagnostics,
    )


async def _run_all(cases: list[dict[str, Any]]) -> list[CaseResult]:
    results: list[CaseResult] = []
    for index, case in enumerate(cases, start=1):
        print(f"\n[{index}/{len(cases)}] RUN {case['id']}")
        result = await _run_case(case)
        status = "PASS" if result.ok else "FAIL"
        print(
            f"[{status}] {result.case_id} returned={result.returned} "
            f"rows={result.output_rows} log={result.log_file} "
            f"output={result.output_file}"
        )
        diag = result.diagnostics or {}
        sources = [
            str(item.get("source") or "")
            for item in diag.get("extracts", [])
            if item.get("source")
        ]
        fields = ((diag.get("output") or {}).get("fields") or [])[:8]
        if sources or fields:
            print(f"  diag sources={sources or []} fields={fields or []}")
        strategies = diag.get("universal_strategies") or {}
        if strategies:
            pre_extract = strategies.get("pre_extract") or {}
            flags: list[str] = []
            if pre_extract.get("triggered"):
                flags.append(
                    "pre_extract="
                    f"{pre_extract.get('last_source')} "
                    f"{pre_extract.get('last_rows')}/{pre_extract.get('last_target')}"
                )
            if (strategies.get("same_page_nav_guard") or {}).get("triggered"):
                flags.append(
                    "same_page_nav_guard="
                    f"{(strategies.get('same_page_nav_guard') or {}).get('count')}"
                )
            if (strategies.get("batch_post_extract_gate") or {}).get("triggered"):
                flags.append(
                    "batch_post_extract_gate="
                    f"{(strategies.get('batch_post_extract_gate') or {}).get('count')}"
                )
            if (strategies.get("extract_click_guard") or {}).get("triggered"):
                flags.append(
                    "extract_click_guard="
                    f"{(strategies.get('extract_click_guard') or {}).get('count')}"
                )
            if (strategies.get("dom_cards") or {}).get("used"):
                flags.append(
                    "dom_cards="
                    f"{(strategies.get('dom_cards') or {}).get('count')}"
                )
            if (strategies.get("full_page_or_ax") or {}).get("used"):
                flags.append(
                    "full_page_or_ax="
                    f"{(strategies.get('full_page_or_ax') or {}).get('count')}"
                )
            if flags:
                print(f"  strategies {'; '.join(flags)}")
        skill_summary = diag.get("skill_summary") or {}
        if skill_summary.get("matched"):
            print(
                "  skills="
                f"{skill_summary.get('skills') or []} "
                f"extracted_by_skill={skill_summary.get('extracted_by_skill')}"
            )
        gate_summary = diag.get("batch_post_extract_gate") or {}
        if gate_summary.get("triggered"):
            print(
                "  batch_post_extract_gate="
                f"count={gate_summary.get('count')} "
                f"next={gate_summary.get('last_next_action')!r} "
                f"rows={gate_summary.get('last_rows')}/{gate_summary.get('last_row_target')} "
                f"pages={gate_summary.get('last_pages')}/{gate_summary.get('last_page_target')}"
            )
        last_done = diag.get("last_done") or {}
        if last_done:
            done_meta = last_done.get("metadata") or {}
            print(
                "  done="
                f"success={last_done.get('success')} "
                f"reason={last_done.get('reason')!r} "
                f"rows={done_meta.get('total_rows')}/{done_meta.get('row_target')} "
                f"target_reached={done_meta.get('target_reached')}"
            )
        if not result.ok and diag.get("failure_type"):
            print(f"  failure_type={diag['failure_type']}")
        if result.error:
            print(f"  error={result.error}")
        results.append(result)
    return results


def _write_report(results: list[CaseResult]) -> Path:
    return _write_report_with_drift(results, snapshot_drift=[])


def _write_report_with_drift(
    results: list[CaseResult],
    snapshot_drift: list[dict[str, Any]],
    skill_replay_check: dict[str, Any] | None = None,
    skill_lint: dict[str, Any] | None = None,
) -> Path:
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    report = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "total": len(results),
        "passed": sum(1 for result in results if result.ok),
        "failed": sum(1 for result in results if not result.ok),
        "skill_summary": _summarize_skill_usage(results),
        "skill_lint": skill_lint or {},
        "skill_replay_check": skill_replay_check or {},
        "snapshot_drift": snapshot_drift,
        "results": [result.__dict__ for result in results],
    }
    path = REPORTS_DIR / f"agent_cases_{datetime.now():%Y%m%d_%H%M%S}.json"
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def _load_snapshot_drift(limit: int) -> list[dict[str, Any]]:
    try:
        return compare_latest_snapshots(DEFAULT_SNAPSHOT_DIR, limit=limit)
    except Exception as exc:  # pragma: no cover - diagnostic fallback only
        return [{"error": repr(exc)}]


def _print_snapshot_drift(snapshot_drift: list[dict[str, Any]]) -> None:
    if not snapshot_drift:
        print("Snapshot drift: no comparable extraction snapshot pairs.")
        return
    print("Snapshot drift:")
    for item in snapshot_drift:
        if item.get("error"):
            print(f"  error={item['error']}")
            continue
        field_delta = item.get("field_delta") or {}
        baseline = item.get("baseline") or {}
        candidate = item.get("candidate") or {}
        print(
            f"  {item.get('group_key','')} match={item.get('match')} "
            f"score={item.get('score')} rows="
            f"{baseline.get('row_count')}->{candidate.get('row_count')} "
            f"fields+={field_delta.get('added', [])} "
            f"fields-={field_delta.get('removed', [])}"
        )


def _print_skill_summary(results: list[CaseResult]) -> None:
    summary = _summarize_skill_usage(results)
    if not summary.get("cases_with_skill"):
        print("Skill summary: no runtime skills matched.")
        return
    print(
        "Skill summary: "
        f"cases={summary['cases_with_skill']} "
        f"extracted={summary['cases_extracted_by_skill']} "
        f"success_rate={summary['skill_success_rate']} "
        f"fallbacks={summary['fallback_count']} "
        f"verify_failures={summary['verification_failure_count']}"
    )
    print(f"  skills_used={summary['skills_used']}")


def _load_skill_replay_check(limit: int) -> dict[str, Any]:
    if int(limit or 0) <= 0:
        return {}
    try:
        return check_skill_replays(limit=int(limit))
    except Exception as exc:  # pragma: no cover - diagnostic fallback only
        return {
            "checked": 0,
            "passed": 0,
            "failed": 1,
            "ok": False,
            "by_skill": {},
            "results": [],
            "error": repr(exc),
        }


def _print_skill_replay_check(summary: dict[str, Any]) -> None:
    if not summary:
        return
    if summary.get("error"):
        print(f"Skill replay check: error={summary['error']}")
        return
    print(
        "Skill replay check: "
        f"ok={summary.get('ok')} "
        f"checked={summary.get('checked')} "
        f"passed={summary.get('passed')} "
        f"failed={summary.get('failed')}"
    )
    for skill_name, item in (summary.get("by_skill") or {}).items():
        print(
            f"  {skill_name}: checked={item.get('checked')} "
            f"passed={item.get('passed')} failed={item.get('failed')}"
        )
    for result in summary.get("results") or []:
        if result.get("ok"):
            continue
        print(
            f"  [FAIL] {result.get('expected_skill')} "
            f"current={result.get('current_first_match')} "
            f"rows={result.get('row_count')}/{result.get('expected_rows')} "
            f"missing={result.get('missing_fields') or []} "
            f"verify_failures={len(result.get('verification_failures') or [])} "
            f"path={result.get('path')}"
        )


def _load_skill_lint(*, include_files: bool = False, strict: bool = False) -> dict[str, Any]:
    try:
        result = lint_skill_registry(include_files=include_files)
    except Exception as exc:  # pragma: no cover - diagnostic fallback only
        return {
            "ok": False,
            "checked": 0,
            "errors": 1,
            "warnings": 0,
            "issues": [],
            "strict": strict,
            "include_files": include_files,
            "error": repr(exc),
        }
    result["strict"] = strict
    result["include_files"] = include_files
    if strict and int(result.get("warnings") or 0):
        result["ok"] = False
    return result


def _print_skill_lint(summary: dict[str, Any]) -> None:
    if not summary:
        return
    if summary.get("error"):
        print(f"Skill lint: error={summary['error']}")
        return
    print(
        "Skill lint: "
        f"ok={summary.get('ok')} "
        f"checked={summary.get('checked')} "
        f"errors={summary.get('errors')} "
        f"warnings={summary.get('warnings')} "
        f"strict={summary.get('strict')} "
        f"include_files={summary.get('include_files')}"
    )
    for issue in summary.get("issues") or []:
        print(
            f"  [{str(issue.get('severity') or '').upper()}] "
            f"{issue.get('skill')} {issue.get('code')}: {issue.get('message')}"
        )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="List or run VSpider generic capability regression cases."
    )
    parser.add_argument("--cases", default=str(DEFAULT_CASES), help="case JSON path")
    parser.add_argument("--id", action="append", default=[], help="case id to run")
    parser.add_argument(
        "--category",
        action="append",
        default=[],
        help="category filter, e.g. interaction or extraction",
    )
    parser.add_argument(
        "--capability",
        action="append",
        default=[],
        help="capability filter, e.g. form, tooltip, table_extract",
    )
    parser.add_argument(
        "--run",
        action="store_true",
        help="actually open browsers and execute cases; omitted means list only",
    )
    parser.add_argument(
        "--include-disabled",
        action="store_true",
        help="include cases marked enabled=false when listing by category/capability/all",
    )
    parser.add_argument(
        "--snapshot-drift-limit",
        type=int,
        default=10,
        help="maximum extraction snapshot drift comparisons to include in run reports",
    )
    parser.add_argument(
        "--skill-replay-check-limit",
        type=int,
        default=0,
        help="offline-check the most recent N skill replay snapshots and fail on drift",
    )
    parser.add_argument(
        "--skill-replay-check-only",
        action="store_true",
        help="run only the skill replay gate without listing or executing cases",
    )
    parser.add_argument(
        "--skip-skill-lint",
        action="store_true",
        help="skip the default runtime skill lint gate before browser runs",
    )
    parser.add_argument(
        "--skill-lint-only",
        action="store_true",
        help="run only the runtime skill lint gate without listing or executing cases",
    )
    parser.add_argument(
        "--skill-lint-include-files",
        action="store_true",
        help="warn about unregistered skill files during the lint gate",
    )
    parser.add_argument(
        "--skill-lint-strict",
        action="store_true",
        help="treat skill lint warnings as failures",
    )
    args = parser.parse_args()

    if args.skill_lint_only:
        skill_lint = _load_skill_lint(
            include_files=bool(args.skill_lint_include_files),
            strict=bool(args.skill_lint_strict),
        )
        _print_skill_lint(skill_lint)
        return 0 if skill_lint.get("ok") else 1

    if args.skill_replay_check_only:
        skill_replay_check = _load_skill_replay_check(max(1, int(args.skill_replay_check_limit or 20)))
        _print_skill_replay_check(skill_replay_check)
        return 0 if skill_replay_check.get("ok") else 1

    cases = _load_cases(Path(args.cases))
    selected = _filter_cases(
        cases,
        set(args.id) or None,
        set(args.category) or None,
        set(args.capability) or None,
        include_disabled=bool(args.include_disabled),
    )
    if not selected:
        print("No cases matched.")
        return 2

    print(f"Selected {len(selected)} case(s):")
    for case in selected:
        _print_case(case)

    if not args.run:
        print("\nList-only mode. Add --run to execute browser regression cases.")
        return 0

    skill_lint = _load_skill_lint(
        include_files=bool(args.skill_lint_include_files),
        strict=bool(args.skill_lint_strict),
    ) if not args.skip_skill_lint else {}
    _print_skill_lint(skill_lint)
    if skill_lint and not skill_lint.get("ok"):
        print("\nStopped before browser run because skill lint failed.")
        return 1

    ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)
    LOGS_DIR.mkdir(parents=True, exist_ok=True)
    results = asyncio.run(_run_all(selected))
    snapshot_drift = _load_snapshot_drift(max(0, int(args.snapshot_drift_limit)))
    skill_replay_check = _load_skill_replay_check(max(0, int(args.skill_replay_check_limit)))
    report = _write_report_with_drift(
        results,
        snapshot_drift=snapshot_drift,
        skill_replay_check=skill_replay_check,
        skill_lint=skill_lint,
    )
    failed = [result for result in results if not result.ok]
    _print_skill_summary(results)
    _print_skill_replay_check(skill_replay_check)
    _print_snapshot_drift(snapshot_drift)
    print(
        f"\nReport: {report}\n"
        f"Passed: {len(results) - len(failed)}/{len(results)}"
    )
    if failed:
        return 1
    if skill_replay_check and not skill_replay_check.get("ok"):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
