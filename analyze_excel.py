import openpyxl
import sys

try:
    wb = openpyxl.load_workbook(r'D:\pythontest\VSpider\workspace\artifacts\output_20260506_120656.xlsx')
    ws = wb.active
    
    print(f'Total rows: {ws.max_row - 1}')
    print(f'Columns: {[c.value for c in ws[1]]}')
    
    print('\n--- First 5 rows ---')
    for r in range(2, min(7, ws.max_row + 1)):
        print([c.value for c in ws[r]])
    
    print('\n--- Last 5 rows ---')
    for r in range(max(2, ws.max_row - 4), ws.max_row + 1):
        print([c.value for c in ws[r]])
    
    urls = [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]
    print(f'\n--- Unique URLs: {len(set(urls))}/{len(urls)} ---')
    
    dupes = [u for u in urls if urls.count(u) > 1]
    print(f'Duplicate URLs: {set(dupes) if dupes else "None"}')
    
except ImportError:
    print("openpyxl not installed, trying pandas...")
    import pandas as pd
    df = pd.read_excel(r'D:\pythontest\VSpider\workspace\artifacts\output_20260506_120656.xlsx')
    print(f'Total rows: {len(df)}')
    print(f'Columns: {list(df.columns)}')
    print('\n--- First 5 rows ---')
    print(df.head().to_string())
    print('\n--- Last 5 rows ---')
    print(df.tail().to_string())
    print(f'\n--- Unique URLs: {df.iloc[:,1].nunique()}/{len(df)} ---')
    dupes = df[df.iloc[:,1].duplicated(keep=False)]
    print(f'Duplicates: {len(dupes)} rows' if len(dupes) > 0 else 'No duplicates')
