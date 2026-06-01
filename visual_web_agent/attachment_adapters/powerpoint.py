"""PowerPoint / ODP presentation adapter.

Returns a short text excerpt of slide titles + body text. Lazy imports
``python-pptx`` / ``odfpy``; falls back to a placeholder when neither is
available.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from visual_web_agent.io_contract.input_contract import AttachmentSpec

from .base import AdapterResult


_DEFAULT_EXCERPT_CHARS = 4000


def adapt(
    spec: AttachmentSpec,
    *,
    goal: str = "",
    options: dict[str, Any] | None = None,
) -> AdapterResult:
    options = options or {}
    max_chars = int(options.get("max_chars") or _DEFAULT_EXCERPT_CHARS)
    path = Path(spec.path)
    if not path.exists():
        return AdapterResult(
            kind="powerpoint", ok=False, reasons=[f"attachment_not_found:{path}"]
        )

    suffix = path.suffix.lower()
    mime = (spec.mime or "").lower()

    if suffix in {".pptx", ".pptm"} or mime in {
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/vnd.ms-powerpoint.presentation.macroEnabled.12",
    }:
        return _read_pptx(path, max_chars=max_chars, spec=spec)

    if suffix == ".ppt" or mime == "application/vnd.ms-powerpoint":
        return _placeholder(
            path, spec,
            reason="legacy_ppt_extractor_not_installed",
            hint="convert .ppt to .pptx via PowerPoint / LibreOffice for full text",
        )

    if suffix == ".odp" or mime == "application/vnd.oasis.opendocument.presentation":
        return _read_odp(path, max_chars=max_chars, spec=spec)

    return _placeholder(path, spec, reason="unsupported_presentation_suffix")


def _read_pptx(path: Path, *, max_chars: int, spec: AttachmentSpec) -> AdapterResult:
    try:
        from pptx import Presentation  # type: ignore
    except Exception:
        return _placeholder(path, spec, reason="python_pptx_unavailable")

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        return AdapterResult(
            kind="powerpoint", ok=False, reasons=[f"pptx_open_failed:{exc}"]
        )

    slides_text: list[str] = []
    slide_count = 0
    for index, slide in enumerate(prs.slides):
        slide_count += 1
        title = ""
        body_parts: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            tf = shape.text_frame
            for paragraph in tf.paragraphs:
                text = "".join((run.text or "") for run in paragraph.runs).strip()
                if not text:
                    continue
                if not title and shape == slide.shapes.title:
                    title = text
                else:
                    body_parts.append(text)
        block_lines = []
        if title:
            block_lines.append(f"# Slide {index + 1}: {title}")
        else:
            block_lines.append(f"# Slide {index + 1}")
        block_lines.extend(body_parts)
        slides_text.append("\n".join(block_lines))
        if sum(len(b) for b in slides_text) >= max_chars:
            break

    body = "\n\n".join(slides_text)
    excerpt = body[:max_chars]
    return AdapterResult(
        kind="powerpoint",
        ok=True,
        text=excerpt,
        metadata={
            "filename": spec.filename or path.name,
            "slide_count": slide_count,
            "total_chars": len(body),
            "excerpt_chars": len(excerpt),
            "source_suffix": path.suffix.lower(),
        },
    )


def _read_odp(path: Path, *, max_chars: int, spec: AttachmentSpec) -> AdapterResult:
    try:
        from odf.opendocument import load as _odf_load  # type: ignore
        from odf import text as _odf_text  # type: ignore
    except Exception:
        return _placeholder(path, spec, reason="odfpy_unavailable")

    try:
        document = _odf_load(str(path))
    except Exception as exc:
        return AdapterResult(
            kind="powerpoint", ok=False, reasons=[f"odp_open_failed:{exc}"]
        )

    chunks: list[str] = []
    for paragraph in document.getElementsByType(_odf_text.P):
        text = _node_text(paragraph)
        if text:
            chunks.append(text)
        if sum(len(c) for c in chunks) >= max_chars:
            break

    body = "\n".join(chunks)
    excerpt = body[:max_chars]
    return AdapterResult(
        kind="powerpoint",
        ok=True,
        text=excerpt,
        metadata={
            "filename": spec.filename or path.name,
            "paragraph_count": len(chunks),
            "total_chars": len(body),
            "excerpt_chars": len(excerpt),
            "source_suffix": ".odp",
        },
    )


def _node_text(node: Any) -> str:
    parts: list[str] = []
    for child in node.childNodes:
        if child.nodeType == 3:  # TEXT_NODE
            parts.append(child.data)
        elif hasattr(child, "childNodes"):
            parts.append(_node_text(child))
    return "".join(parts).strip()


def _placeholder(
    path: Path,
    spec: AttachmentSpec,
    *,
    reason: str,
    hint: str = "",
) -> AdapterResult:
    caption = f"[presentation attachment: {spec.filename or path.name}, {reason}]"
    reasons = [reason]
    if hint:
        reasons.append(hint)
    return AdapterResult(
        kind="powerpoint",
        ok=True,
        text=caption,
        metadata={
            "filename": spec.filename or path.name,
            "size": spec.size or (path.stat().st_size if path.exists() else 0),
        },
        reasons=reasons,
    )
