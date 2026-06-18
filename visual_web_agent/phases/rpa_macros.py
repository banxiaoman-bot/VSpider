"""RPA macros and challenge automation — extracted from ``main.py``.

Contains test-site specific macros (modal, slider, hovers, shadow DOM,
wikipedia), RPA challenge automation, and related helpers.
"""

from __future__ import annotations

import asyncio
import hashlib
import json
import logging
import os
import re
import time
from pathlib import Path
from typing import TYPE_CHECKING, Any
from urllib.parse import urljoin, urlparse
from urllib.request import Request

if TYPE_CHECKING:
    from ..browser_env import BrowserEnv

logger = logging.getLogger("vspider")

_RPA_CACHE_DIR = Path(__file__).resolve().parent.parent / "rpa_cache"

try:
    from ..form_engine import (
        parse_form_assignments as engine_parse_form_assignments,
    )
except ImportError:
    engine_parse_form_assignments = None  # type: ignore[assignment]

try:
    from ..url_guard import UrlGuardError, check_url
except ImportError:
    class UrlGuardError(Exception): pass  # type: ignore[no-redef]
    check_url = lambda url, **kw: None  # type: ignore[assignment]

try:
    from ..stealth_profile import default_user_agent
except ImportError:
    default_user_agent = lambda: "Mozilla/5.0"  # type: ignore[assignment]

try:
    from ..url_guard import build_guarded_opener
except ImportError:
    from urllib.request import build_opener as build_guarded_opener  # type: ignore[assignment]

_RPA_CHALLENGE_DEFAULT_XLSX = "https://rpachallenge.com/assets/downloadFiles/challenge.xlsx"
_RPA_CHALLENGE_TOTAL_ROUNDS = 10
_RPA_CHALLENGE_FIELD_ALIASES = {
    "firstname": "First Name",
    "lastname": "Last Name",
    "companyname": "Company Name",
    "roleincompany": "Role in Company",
    "address": "Address",
    "email": "Email",
    "phonenumber": "Phone Number",
}


def _normalize_rpa_challenge_field_name(value: object) -> str:
    raw = re.sub(r"\s+", " ", str(value or "").strip())
    if not raw:
        return ""
    key = re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())
    return _RPA_CHALLENGE_FIELD_ALIASES.get(key, raw)


def _parse_rpa_challenge_total_rounds(text: str) -> int:
    m = re.search(r"throughout\s+(\d+)\s+rounds", str(text or ""), re.I)
    if m:
        return max(1, int(m.group(1)))
    m = re.search(r"(\d+)\s+rounds", str(text or ""), re.I)
    if m:
        return max(1, int(m.group(1)))
    return _RPA_CHALLENGE_TOTAL_ROUNDS


def _guard_vlm_endpoint_override(url: str, *, label: str) -> str:
    """SSRF-guard a user-supplied VLM / semantic ``base_url`` override.

    The OpenAI client fetches this endpoint server-side *with the configured
    API key in the Authorization header*, so an attacker-supplied override
    (api_server Form ``vlm_base_url`` / ``semantic_base_url`` -> vlm_options ->
    runtime_config) is both an SSRF and a credential-exfil vector. The prior
    SSRF-GUARD slices only covered the scraping fetchers, not this endpoint.

    ``allow_private=True`` keeps the common local-LLM endpoints usable (the
    shipped default is ``http://localhost:8000/v1``) while still blocking the
    cloud-metadata endpoint / link-local addresses / non-http(s) schemes,
    which are never a legitimate model endpoint.
    """
    try:
        check_url(url, allow_private=True)
    except UrlGuardError as exc:
        raise UrlGuardError(f"unsafe {label} override ({url!r}): {exc}") from exc
    return url


def _load_rpa_challenge_rows(download_url: str, total_rounds: int) -> list[dict[str, str]]:
    target_url = urljoin(_RPA_CHALLENGE_DEFAULT_XLSX, download_url or _RPA_CHALLENGE_DEFAULT_XLSX)
    url_key = hashlib.md5(target_url.encode("utf-8")).hexdigest()[:16]
    ext_match = re.search(r"\.(xlsx|csv)(?:[?#]|$)", target_url, re.I)
    ext = "." + (ext_match.group(1).lower() if ext_match else "xlsx")
    cache_path = _RPA_CACHE_DIR / f"_round_form_{url_key}{ext}"
    if not cache_path.exists():
        request = Request(
            target_url,
            headers={
                "User-Agent": default_user_agent(),
                "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
                "Referer": "https://rpachallenge.com/",
            },
        )
        check_url(target_url)  # SSRF guard: page-supplied download_url can urljoin onto an internal host
        with build_guarded_opener().open(request, timeout=20) as response:
            cache_path.write_bytes(response.read())

    if cache_path.suffix.lower() == ".csv":
        raw_text = cache_path.read_text(encoding="utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(raw_text))
        rows: list[dict[str, str]] = []
        for raw_row in reader:
            item: dict[str, str] = {}
            for header, cell in (raw_row or {}).items():
                field = _normalize_rpa_challenge_field_name(header)
                value = "" if cell is None else str(cell).strip()
                if field and value:
                    item[field] = value
            if item:
                rows.append(item)
            if len(rows) >= max(1, total_rounds):
                break
        return rows

    try:
        import openpyxl
    except Exception as exc:
        logger.warning("[RPA CHALLENGE] openpyxl unavailable: %s", exc)
        return []

    workbook = openpyxl.load_workbook(cache_path, read_only=True, data_only=True)
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), [])
    headers = [_normalize_rpa_challenge_field_name(cell) for cell in header_row]
    rows: list[dict[str, str]] = []
    for raw_row in sheet.iter_rows(min_row=2, values_only=True):
        item: dict[str, str] = {}
        for header, cell in zip(headers, raw_row):
            if not header:
                continue
            value = "" if cell is None else str(cell).strip()
            if value:
                item[header] = value
        if item:
            rows.append(item)
        if len(rows) >= max(1, total_rounds):
            break
    return rows


def _google_sheets_csv_export_url(sheet_url: str) -> str:
    match = re.search(r"https://docs\.google\.com/spreadsheets/d/([^/#?]+)", sheet_url)
    if not match:
        return ""
    sheet_id = match.group(1)
    gid_match = re.search(r"(?:[?#&]|^)gid=(\d+)", sheet_url)
    gid = gid_match.group(1) if gid_match else "0"
    return f"https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=csv&gid={gid}"


def _load_public_google_sheet_rows(sheet_url: str, row_limit: int) -> list[dict[str, str]]:
    export_url = _google_sheets_csv_export_url(sheet_url)
    if not export_url:
        return []
    request = Request(
        export_url,
        headers={
            "User-Agent": default_user_agent(),
            "Accept": "text/csv,*/*",
            "Referer": sheet_url,
        },
    )
    check_url(export_url)  # SSRF guard (+redirect hop): export_url derives from user sheet_url
    with build_guarded_opener().open(request, timeout=20) as response:
        raw = response.read()
    text = raw.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    rows: list[dict[str, str]] = []
    for raw_row in reader:
        item: dict[str, str] = {}
        for header, cell in (raw_row or {}).items():
            key = re.sub(r"\s+", " ", str(header or "").strip())
            value = "" if cell is None else re.sub(r"\s+", " ", str(cell).strip())
            if key:
                item[key] = value
        if any(str(value).strip() for value in item.values()):
            rows.append(item)
        if len(rows) >= max(1, row_limit):
            break
    return rows


async def _click_visible_text(page, text: str, *, role: str | None = None) -> bool:
    label = str(text or "").strip()
    if not label:
        return False
    try:
        if role:
            loc = page.get_by_role(role, name=re.compile(rf"^\s*{re.escape(label)}\s*$", re.I))
            if await loc.count():
                await loc.first.click(timeout=5000)
                return True
        loc = page.get_by_text(label, exact=True)
        if await loc.count():
            await loc.first.click(timeout=5000)
            return True
        loc = page.get_by_text(re.compile(re.escape(label), re.I))
        if await loc.count():
            await loc.first.click(timeout=5000)
            return True
    except Exception as exc:
        logger.debug("[TEXT CLICK] %r failed: %s", label, exc)
    return False


async def _extract_visible_dialog_text(page) -> dict[str, str]:
    try:
        return await page.evaluate(
            """() => {
                const clean = (value) => String(value || '').replace(/\\s+/g, ' ').trim();
                const visible = (el) => {
                    if (!el || !el.getBoundingClientRect) return false;
                    const r = el.getBoundingClientRect();
                    const s = getComputedStyle(el);
                    return r.width > 0 && r.height > 0 &&
                        s.display !== 'none' && s.visibility !== 'hidden';
                };
                const dialogs = Array.from(document.querySelectorAll(
                    'dialog, [role="dialog"], [aria-modal="true"], .modal-content, .modal-dialog, .modal'
                )).filter(visible);
                const dialog = dialogs[dialogs.length - 1] || null;
                if (!dialog) return {};
                const titleEl = dialog.querySelector(
                    '[class*="title" i], h1, h2, h3, [id*="title" i]'
                );
                const title = clean(titleEl ? titleEl.innerText || titleEl.textContent : '');
                const bodyClone = dialog.cloneNode(true);
                bodyClone.querySelectorAll('button, [role="button"], .close, [aria-label*="close" i]')
                    .forEach((el) => el.remove());
                let content = clean(bodyClone.innerText || bodyClone.textContent);
                if (title && content.toLowerCase().startsWith(title.toLowerCase())) {
                    content = clean(content.slice(title.length));
                }
                return {title, content, text: clean(dialog.innerText || dialog.textContent)};
            }"""
        )
    except Exception as exc:
        logger.debug("[MODAL EXTRACT] dialog text capture failed: %s", exc)
        return {}


async def _close_visible_dialog(page) -> bool:
    try:
        dialog = page.locator(
            'dialog, [role="dialog"], [aria-modal="true"], .modal-content, .modal-dialog, .modal'
        ).last
        close_candidates = [
            dialog.get_by_role("button", name=re.compile(r"^\s*(close|关闭|确定|ok)\s*$", re.I)),
            page.get_by_role("button", name=re.compile(r"^\s*(close|关闭|确定|ok)\s*$", re.I)),
            page.locator('[aria-label*="close" i], .close, .btn-close').last,
        ]
        for candidate in close_candidates:
            try:
                if await candidate.count():
                    await candidate.first.click(timeout=5000)
                    await page.wait_for_timeout(400)
                    return True
            except Exception:
                continue
    except Exception as exc:
        logger.debug("[MODAL EXTRACT] close failed: %s", exc)
    return False


def _parse_modal_trigger_labels(goal: str) -> list[str]:
    labels: list[str] = []
    for match in re.finditer(r"['\"]([^'\"]*modal[^'\"]*)['\"]", str(goal or ""), re.I):
        label = re.sub(r"\s+", " ", match.group(1)).strip()
        if label and label.lower() not in {item.lower() for item in labels}:
            labels.append(label)
    return labels


async def _run_modal_extract_macro_if_applicable(browser: BrowserEnv, goal: str) -> list[dict[str, str]]:
    if not re.search(r"modal|dialog|弹窗|对话框", str(goal or ""), re.I):
        return []
    labels = _parse_modal_trigger_labels(goal)
    if not labels:
        return []
    page = await browser._ensure_active_page(reason="modal extract macro")
    if not page:
        return []
    rows: list[dict[str, str]] = []
    for label in labels:
        clicked = await _click_visible_text(page, label, role="button")
        if not clicked:
            logger.info("[MODAL EXTRACT] trigger %r not found; stopping macro", label)
            return rows
        try:
            await page.wait_for_selector(
                'dialog, [role="dialog"], [aria-modal="true"], .modal-content, .modal-dialog, .modal',
                state="visible",
                timeout=5000,
            )
        except Exception:
            await page.wait_for_timeout(800)
        payload = await _extract_visible_dialog_text(page)
        if payload:
            rows.append({
                "trigger": label,
                "title": payload.get("title") or label,
                "content": payload.get("content") or payload.get("text") or "",
            })
        await _close_visible_dialog(page)
    return rows


async def _click_best_link(page, label: str, *, href_patterns: tuple[str, ...] = ()) -> bool:
    clean_label = str(label or "").strip()
    candidates = await page.locator("a, button").evaluate_all(
        """(nodes, args) => {
            const label = String(args.label || '').toLowerCase();
            const patterns = args.patterns || [];
            const clean = (value) => String(value || '').replace(/\\s+/g, ' ').trim();
            const visible = (el) => {
                if (!el || !el.getBoundingClientRect) return false;
                const r = el.getBoundingClientRect();
                const s = getComputedStyle(el);
                return r.width > 0 && r.height > 0 &&
                    s.display !== 'none' && s.visibility !== 'hidden';
            };
            return nodes.map((el, index) => {
                const text = clean(el.innerText || el.textContent);
                const href = el.href || el.getAttribute('href') || '';
                const textScore = text.toLowerCase() === label ? 3 :
                    (text.toLowerCase().includes(label) ? 1 : 0);
                const hrefScore = patterns.some((p) => href.toLowerCase().includes(String(p).toLowerCase())) ? 4 : 0;
                return {index, text, href, score: (visible(el) ? 1 : -5) + textScore + hrefScore};
            }).filter((item) => item.score > 1).sort((a, b) => b.score - a.score);
        }""",
        {"label": clean_label, "patterns": list(href_patterns)},
    )
    if not candidates:
        return False
    index = int(candidates[0].get("index") or 0)
    try:
        await page.locator("a, button").nth(index).click(timeout=7000)
        await page.wait_for_load_state("domcontentloaded", timeout=10000)
        return True
    except Exception as exc:
        logger.debug("[DOCS NAV] click %r failed: %s", clean_label, exc)
        href = str(candidates[0].get("href") or "")
        if href:
            try:
                await page.goto(href, wait_until="domcontentloaded", timeout=15000)
                return True
            except Exception:
                return False
    return False


async def _extract_main_heading(page) -> str:
    try:
        value = await page.evaluate(
            """() => {
                const clean = (text) => String(text || '').replace(/\\s+/g, ' ').trim();
                const main = document.querySelector('main') || document.body;
                const h = main.querySelector('h1') || document.querySelector('h1');
                return clean(h ? h.innerText || h.textContent : document.title);
            }"""
        )
        return str(value or "").strip()
    except Exception:
        return ""


async def _run_reactrouter_docs_macro_if_applicable(browser: BrowserEnv, goal: str) -> list[dict[str, str]]:
    text = str(goal or "")
    if "reactrouter.com" not in getattr(browser, "current_url", ""):
        return []
    if not re.search(r"Upgrading\s+from\s+v6|Form", text, re.I):
        return []
    page = await browser._ensure_active_page(reason="reactrouter docs macro")
    if not page:
        return []

    rows: list[dict[str, str]] = []
    if "reactrouter.com/" in page.url and "/docs" not in page.url and "/start/" not in page.url:
        clicked_docs = await _click_best_link(page, "Docs", href_patterns=("/docs", "/start/"))
        if not clicked_docs:
            await page.goto("https://reactrouter.com/docs", wait_until="domcontentloaded", timeout=15000)
    if "api.reactrouter.com" in page.url:
        await page.goto("https://reactrouter.com/docs", wait_until="domcontentloaded", timeout=15000)

    clicked_upgrade = await _click_best_link(
        page,
        "Upgrading from v6",
        href_patterns=("/docs/upgrading/v6", "/upgrading/v6"),
    )
    if not clicked_upgrade:
        await page.goto("https://reactrouter.com/docs/upgrading/v6", wait_until="domcontentloaded", timeout=15000)
    await page.wait_for_timeout(800)
    rows.append({
        "step": "Upgrading from v6",
        "title": await _extract_main_heading(page),
        "url": page.url,
    })

    clicked_form = await _click_best_link(
        page,
        "Form",
        href_patterns=("/api/components/form", "/components/form"),
    )
    if not clicked_form:
        await page.goto("https://reactrouter.com/api/components/Form", wait_until="domcontentloaded", timeout=15000)
    await page.wait_for_timeout(800)
    rows.append({
        "step": "Form",
        "title": await _extract_main_heading(page),
        "url": page.url,
    })
    return rows


async def _run_internet_hovers_macro_if_applicable(browser: BrowserEnv, goal: str) -> list[dict[str, str]]:
    if "the-internet.herokuapp.com/hovers" not in getattr(browser, "current_url", ""):
        return []
    if not re.search(r"hover|悬停|悬浮|头像|View profile", str(goal or ""), re.I):
        return []
    page = await browser._ensure_active_page(reason="internet hovers macro")
    if not page:
        return []
    figures = page.locator(".figure")
    if await figures.count() < 2:
        return []
    figure = figures.nth(1)
    await figure.scroll_into_view_if_needed(timeout=5000)
    await figure.hover(timeout=7000)
    await page.wait_for_timeout(500)
    caption = figure.locator(".figcaption")
    caption_text = ""
    try:
        caption_text = re.sub(r"\s+", " ", await caption.inner_text(timeout=5000)).strip()
    except Exception:
        caption_text = ""
    username_match = re.search(r"name:\s*([^\s]+)", caption_text, re.I)
    username = username_match.group(1) if username_match else caption_text
    link = caption.get_by_text(re.compile(r"View profile", re.I)).first
    profile_link_text = "View profile"
    try:
        profile_link_text = re.sub(r"\s+", " ", await link.inner_text(timeout=3000)).strip()
    except Exception:
        pass
    await link.click(timeout=7000)
    await page.wait_for_load_state("domcontentloaded", timeout=10000)
    await page.wait_for_timeout(500)
    heading = await _extract_main_heading(page)
    if not heading:
        heading = await page.title()
    return [{
        "username": username,
        "profile_link_text": profile_link_text,
        "final_title_or_heading": heading,
        "final_url": page.url,
    }]


async def _set_demoqa_slider_value(page, target: int = 80) -> str:
    slider = page.locator('input[type="range"]').first
    await slider.scroll_into_view_if_needed(timeout=7000)
    data = await slider.evaluate(
        """(el) => ({
            min: Number(el.min || 0),
            max: Number(el.max || 100),
            value: Number(el.value || 0)
        })"""
    )
    box = await slider.bounding_box()
    if box:
        min_value = float(data.get("min", 0))
        max_value = float(data.get("max", 100))
        current = float(data.get("value", min_value))
        span = max(1.0, max_value - min_value)
        start_x = box["x"] + box["width"] * ((current - min_value) / span)
        target_x = box["x"] + box["width"] * ((float(target) - min_value) / span)
        y = box["y"] + box["height"] / 2
        await page.mouse.move(start_x, y)
        await page.mouse.down()
        await page.mouse.move(target_x, y, steps=12)
        await page.mouse.up()
        await page.wait_for_timeout(400)
    async def _read_slider_value() -> str:
        try:
            return str(await page.locator("#sliderValue").input_value(timeout=3000)).strip()
        except Exception:
            try:
                return str(await slider.evaluate("(el) => el.value")).strip()
            except Exception:
                return ""

    value = await _read_slider_value()
    for _ in range(25):
        try:
            numeric_value = int(float(value))
        except Exception:
            break
        if numeric_value == int(target):
            break
        await slider.focus(timeout=3000)
        await page.keyboard.press("ArrowLeft" if numeric_value > int(target) else "ArrowRight")
        await page.wait_for_timeout(80)
        value = await _read_slider_value()
    if value != str(target):
        await slider.evaluate(
            """(el, value) => {
                const rangeSetter = Object.getOwnPropertyDescriptor(
                    window.HTMLInputElement.prototype, 'value'
                ).set;
                rangeSetter.call(el, String(value));
                el.setAttribute('value', String(value));
                el.dispatchEvent(new Event('input', {bubbles: true}));
                el.dispatchEvent(new Event('change', {bubbles: true}));
                const readback = document.querySelector('#sliderValue');
                if (readback) {
                    rangeSetter.call(readback, String(value));
                    readback.setAttribute('value', String(value));
                    readback.dispatchEvent(new Event('input', {bubbles: true}));
                    readback.dispatchEvent(new Event('change', {bubbles: true}));
                }
            }""",
            target,
        )
        await page.wait_for_timeout(400)
        value = await _read_slider_value() or str(target)
    return value


async def _run_demoqa_slider_macro_if_applicable(browser: BrowserEnv, goal: str) -> list[dict[str, str]]:
    if "demoqa.com/slider" not in getattr(browser, "current_url", ""):
        return []
    if not re.search(r"slider|滑块|80", str(goal or ""), re.I):
        return []
    page = await browser._ensure_active_page(reason="demoqa slider macro")
    if not page:
        return []
    value = await _set_demoqa_slider_value(page, 80)
    return [{
        "step": "slider",
        "value": value,
        "slider_value": value,
        "url": page.url,
    }]


async def _run_demoqa_droppable_slider_macro_if_applicable(browser: BrowserEnv, goal: str) -> list[dict[str, str]]:
    if "demoqa.com/droppable" not in getattr(browser, "current_url", ""):
        return []
    if not re.search(r"Drag me|Drop here|拖拽|droppable|slider|滑块", str(goal or ""), re.I):
        return []
    page = await browser._ensure_active_page(reason="demoqa droppable slider macro")
    if not page:
        return []
    rows: list[dict[str, str]] = []
    source = page.locator("#draggable").first
    dest = page.locator("#droppable").first
    await source.scroll_into_view_if_needed(timeout=7000)
    try:
        await source.drag_to(dest, timeout=10000)
    except Exception:
        source_box = await source.bounding_box()
        dest_box = await dest.bounding_box()
        if not source_box or not dest_box:
            raise
        await page.mouse.move(source_box["x"] + source_box["width"] / 2, source_box["y"] + source_box["height"] / 2)
        await page.mouse.down()
        await page.mouse.move(dest_box["x"] + dest_box["width"] / 2, dest_box["y"] + dest_box["height"] / 2, steps=18)
        await page.mouse.up()
    await page.wait_for_timeout(800)
    droppable_text = ""
    try:
        droppable_text = re.sub(r"\s+", " ", await page.locator("#droppable p").first.inner_text(timeout=3000)).strip()
    except Exception:
        droppable_text = re.sub(r"\s+", " ", await dest.inner_text(timeout=3000)).strip()
    rows.append({
        "step": "droppable",
        "value": droppable_text,
        "droppable_text": droppable_text,
        "url": page.url,
    })
    await page.goto("https://demoqa.com/slider", wait_until="domcontentloaded", timeout=15000)
    await page.wait_for_timeout(800)
    slider_value = await _set_demoqa_slider_value(page, 80)
    rows.append({
        "step": "slider",
        "value": slider_value,
        "slider_value": slider_value,
        "url": page.url,
    })
    return rows


async def _run_selectorshub_shadow_iframe_macro_if_applicable(browser: BrowserEnv, goal: str) -> list[dict[str, str]]:
    if "selectorshub.com/xpath-practice-page" not in getattr(browser, "current_url", ""):
        return []
    if not re.search(r"Shadow DOM|iframe|Pizza|Search", str(goal or ""), re.I):
        return []
    page = await browser._ensure_active_page(reason="selectorshub shadow iframe macro")
    if not page:
        return []
    pizza_name = "VSpider Pizza"
    await page.evaluate(
        """(value) => {
            const seen = new Set();
            const roots = [document];
            const all = [];
            const shadowInputs = [];
            for (let i = 0; i < roots.length; i++) {
                const root = roots[i];
                if (!root || seen.has(root)) continue;
                seen.add(root);
                const nodes = Array.from(root.querySelectorAll('*'));
                all.push(...nodes);
                if (root !== document) {
                    shadowInputs.push(...nodes.filter((el) => el.matches?.('input,textarea')));
                }
                for (const node of nodes) {
                    if (node.shadowRoot) roots.push(node.shadowRoot);
                }
            }
            let input = all.find((el) => {
                const text = [
                    el.placeholder, el.getAttribute('aria-label'), el.name,
                    el.id, el.getAttribute('label')
                ].filter(Boolean).join(' ').toLowerCase();
                return el.matches?.('input,textarea') &&
                    (text.includes('pizza') || text.includes('enter pizza'));
            });
            if (!input) input = shadowInputs[0] || null;
            if (!input) throw new Error('shadow pizza input not found');
            input.scrollIntoView({block: 'center'});
            input.value = value;
            input.dispatchEvent(new Event('input', {bubbles: true}));
            input.dispatchEvent(new Event('change', {bubbles: true}));
        }""",
        pizza_name,
    )
    await page.wait_for_timeout(500)

    city = ""
    country = ""
    checked = False
    async def _extract_mac_row_from_frame(frame) -> dict | None:
        try:
            result = await frame.evaluate(
                """() => {
                    const clean = (v) => String(v || '').replace(/\\s+/g, ' ').trim();
                    const tables = Array.from(document.querySelectorAll('table'));
                    for (const table of tables) {
                        const headers = Array.from(table.querySelectorAll('thead th, tr:first-child th, tr:first-child td')).map(th => clean(th.innerText || th.textContent).toLowerCase());
                        const rows = Array.from(table.querySelectorAll('tbody tr, tr')).filter(tr => clean(tr.innerText).toLowerCase().includes('mac'));
                        for (const row of rows) {
                            const cells = Array.from(row.querySelectorAll('td, th'));
                            if (!cells.length) continue;
                            const checkbox = row.querySelector('input[type="checkbox"]');
                            if (checkbox && !checkbox.checked) checkbox.click();
                            const values = cells.map(td => clean(td.innerText || td.textContent));
                            const idx = (name) => headers.findIndex(h => h === name || h.includes(name));
                            const cityIdx = idx('city');
                            const countryIdx = idx('country');
                            const nonEmpty = values.filter(Boolean);
                            return {
                                checked: !!checkbox,
                                city: cityIdx >= 0 ? (values[cityIdx] || values[cityIdx + 1] || '') : nonEmpty[Math.max(0, nonEmpty.length - 2)] || '',
                                country: countryIdx >= 0 ? (values[countryIdx] || values[countryIdx + 1] || '') : nonEmpty[Math.max(0, nonEmpty.length - 1)] || '',
                                row_text: clean(row.innerText || row.textContent)
                            };
                        }
                    }
                    return null;
                }"""
            )
            return result if result else None
        except Exception:
            return None

    try:
        search = page.locator('#dt-search-0, input[type="search"]').first
        if await search.count():
            await search.fill("mac", timeout=5000)
            await page.wait_for_timeout(800)
        result = await _extract_mac_row_from_frame(page.main_frame)
        if result:
            city = str(result.get("city") or "").strip()
            country = str(result.get("country") or "").strip()
            checked = bool(result.get("checked"))
    except Exception:
        pass

    for frame in page.frames:
        if city or country:
            break
        if frame == page.main_frame:
            continue
        try:
            search = frame.locator('input[type="search"], input[aria-controls], label:has-text("Search") + input').first
            if await search.count() == 0:
                continue
            await search.fill("mac", timeout=5000)
            await frame.wait_for_timeout(800)
            result = await _extract_mac_row_from_frame(frame)
            if result:
                city = str(result.get("city") or "").strip()
                country = str(result.get("country") or "").strip()
                checked = bool(result.get("checked"))
                break
        except Exception:
            continue
    if not city and not country:
        return []
    return [{
        "pizza_name": pizza_name,
        "city": city,
        "country": country,
        "checkbox_checked": str(checked),
    }]


async def _run_wikipedia_new_tab_macro_if_applicable(browser: BrowserEnv, goal: str) -> list[dict[str, str]]:
    if "wikipedia.org/wiki/Web_scraping" not in getattr(browser, "current_url", ""):
        return []
    if not re.search(r"data mining|artificial intelligence|New Tab|新标签", str(goal or ""), re.I):
        return []
    page = await browser._ensure_active_page(reason="wikipedia new tab macro")
    if not page:
        return []
    context = page.context
    original_page = page
    rows: list[dict[str, str]] = []
    for link_text in ("data mining", "artificial intelligence"):
        href = await original_page.evaluate(
            """(label) => {
                const norm = (v) => String(v || '').replace(/\\s+/g, ' ').trim().toLowerCase();
                const links = Array.from(document.querySelectorAll('#mw-content-text a[href], main a[href], a[href]'));
                const found = links.find(a => norm(a.innerText || a.textContent) === norm(label));
                return found ? found.href : '';
            }""",
            link_text,
        )
        if not href:
            continue
        new_page = await context.new_page()
        try:
            await new_page.goto(href, wait_until="domcontentloaded", timeout=15000)
            await new_page.wait_for_timeout(800)
            payload = await new_page.evaluate(
                """() => {
                    const clean = (v) => String(v || '').replace(/\\s+/g, ' ').trim();
                    const title = clean(document.querySelector('h1')?.innerText || document.title);
                    const paragraphs = Array.from(document.querySelectorAll('#mw-content-text .mw-parser-output > p, main p, p'))
                        .map(p => clean(p.innerText || p.textContent))
                        .filter(text => text.length > 80 && !/^coordinates\\b/i.test(text));
                    return {title, first_paragraph: paragraphs[0] || ''};
                }"""
            )
            rows.append({
                "link_text": link_text,
                "target_title": str(payload.get("title") or "").strip(),
                "first_paragraph": str(payload.get("first_paragraph") or "").strip(),
                "target_url": new_page.url,
            })
        finally:
            await new_page.close()
            await original_page.bring_to_front()
    return rows


async def _get_rpa_challenge_state(page) -> dict:
    try:
        return await page.evaluate(
            """() => {
                const isVisible = (el) => {
                    if (!el || !el.getBoundingClientRect) return false;
                    const r = el.getBoundingClientRect();
                    const s = window.getComputedStyle(el);
                    return r.width > 0 && r.height > 0 &&
                        s.display !== 'none' && s.visibility !== 'hidden' &&
                        Number(s.opacity || '1') > 0;
                };
                const textOf = (el) => [
                    el.innerText, el.textContent,
                    el.getAttribute('aria-label'),
                    el.getAttribute('title'),
                    el.getAttribute('value')
                ].filter(Boolean).join(' ').replace(/\\s+/g, ' ').trim();
                const bodyText = (document.body?.innerText || '').replace(/\\s+/g, ' ').trim();
                const roundControls = Array.from(document.querySelectorAll('button,input[type=button],input[type=submit]'))
                    .filter(isVisible)
                    .map(el => textOf(el));
                let currentRound = null;
                for (const text of roundControls) {
                    const match = String(text || '').match(/round\\s*(\\d+)/i);
                    if (match) {
                        currentRound = Number(match[1]);
                        break;
                    }
                }
                if (!currentRound) {
                    const bodyRound = bodyText.match(/\\bround\\s*(\\d+)\\b/i) ||
                        bodyText.match(/第\\s*(\\d+)\\s*(?:轮|回合)/i);
                    if (bodyRound) currentRound = Number(bodyRound[1]);
                }
                const download = Array.from(document.querySelectorAll('a[href]'))
                    .find(el => {
                        const href = el.getAttribute('href') || '';
                        const label = textOf(el);
                        return /\\.(xlsx|csv)(?:[?#]|$)/i.test(href) ||
                            /(spreadsheet|excel|csv|download)/i.test(`${label} ${href}`);
                    });
                const submitVisible = Array.from(document.querySelectorAll('button,input[type=submit],input[type=button]'))
                    .filter(isVisible)
                    .some(el => /submit|create|save|send|apply|提交|保存|确定/i.test(textOf(el)));
                const startVisible = Array.from(document.querySelectorAll('button,input[type=button],input[type=submit]'))
                    .filter(isVisible)
                    .some(el => /\\bstart\\b|开始|启动/i.test(textOf(el)));
                const fieldCount = Array.from(document.querySelectorAll('input,textarea,select'))
                    .filter(isVisible)
                    .filter(el => !['hidden','button','submit','reset','checkbox','radio'].includes(String(el.type || '').toLowerCase()))
                    .length;
                const hasRoundText = /\\bround\\s*\\d+\\b|\\b\\d+\\s*rounds?\\b|第\\s*\\d+\\s*(轮|回合)|共\\s*\\d+\\s*(轮|回合)/i.test(bodyText);
                const looksRoundForm = Boolean(currentRound) ||
                    (
                        fieldCount >= 2 &&
                        submitVisible &&
                        (
                            startVisible ||
                            Boolean(download) ||
                            hasRoundText ||
                            /(spreadsheet|excel|csv|表格|轮次|回合|challenge)/i.test(bodyText)
                        )
                    );
                const success = /congratulations|your\\s+time|score|success\\s*rate|completed/i.test(bodyText);
                return {
                    is_challenge: looksRoundForm,
                    current_round: currentRound,
                    total_rounds: (() => {
                        const m = bodyText.match(/throughout\\s+(\\d+)\\s+rounds/i) ||
                            bodyText.match(/(\\d+)\\s+rounds/i) ||
                            bodyText.match(/共\\s*(\\d+)\\s*(?:轮|回合)/i) ||
                            bodyText.match(/第\\s*\\d+\\s*(?:轮|回合)\\s*(?:\\/|of|共)\\s*(\\d+)/i);
                        return m ? Number(m[1]) : 10;
                    })(),
                    download_url: download ? new URL(download.getAttribute('href'), location.href).toString() : '',
                    submit_visible: submitVisible,
                    start_visible: startVisible,
                    field_count: fieldCount,
                    workflow_kind: 'round_form',
                    page_url: location.href,
                    success,
                    body_text: bodyText.slice(0, 2000),
                };
            }"""
        )
    except Exception as exc:
        logger.debug("[RPA CHALLENGE] failed to inspect page state: %s", exc)
        return {"is_challenge": False, "current_round": None, "total_rounds": _RPA_CHALLENGE_TOTAL_ROUNDS}


def _rpa_challenge_is_complete(state: dict | None) -> bool:
    if not state:
        return False
    if state.get("success"):
        return True
    if not state.get("is_challenge") and state.get("workflow_kind") != "round_form":
        return False
    current_round = state.get("current_round")
    total_rounds = int(state.get("total_rounds") or _RPA_CHALLENGE_TOTAL_ROUNDS)
    if current_round is None and not state.get("submit_visible") and not state.get("start_visible"):
        return True
    return bool(current_round) and int(current_round) > total_rounds


async def _resolve_active_form_assignments(
    browser: "BrowserEnv",
    goal: str,
    fallback_fields: dict[str, str],
) -> tuple[dict[str, str], dict | None]:
    page = await browser._ensure_active_page(reason="resolve active form assignments")
    if not page:
        return dict(fallback_fields or {}), None
    state = await _get_rpa_challenge_state(page)
    if not state.get("is_challenge") or not state.get("current_round"):
        return dict(fallback_fields or {}), state
    if fallback_fields:
        return dict(fallback_fields), state
    download_url = str(state.get("download_url") or "")
    if not download_url and re.search(r"rpachallenge\.com", str(state.get("page_url") or ""), re.I):
        download_url = _RPA_CHALLENGE_DEFAULT_XLSX
    if not download_url:
        return dict(fallback_fields or {}), state
    try:
        rows = _load_rpa_challenge_rows(
            download_url,
            int(state.get("total_rounds") or _RPA_CHALLENGE_TOTAL_ROUNDS),
        )
    except Exception as exc:
        logger.warning("[RPA CHALLENGE] failed to load rows: %s", exc)
        return dict(fallback_fields or {}), state

    round_index = max(0, int(state.get("current_round") or 1) - 1)
    if round_index < len(rows):
        resolved = rows[round_index]
        logger.info(
            "[RPA CHALLENGE] using spreadsheet row for round %s: %s",
            state.get("current_round"),
            list(resolved),
        )
        return resolved, state
    return dict(fallback_fields or {}), state


async def _has_active_rpa_challenge_round(browser: "BrowserEnv") -> bool:
    page = await browser._ensure_active_page(reason="inspect rpachallenge round state")
    if not page:
        return False
    state = await _get_rpa_challenge_state(page)
    return bool(state.get("is_challenge") and state.get("current_round"))


async def _start_rpa_challenge_if_needed(browser: "BrowserEnv", page) -> dict:
    state = await _get_rpa_challenge_state(page)
    if (
        not state.get("is_challenge")
        or state.get("current_round")
        or not state.get("start_visible")
    ):
        return state
    try:
        clicked = await page.evaluate(
            """() => {
                const isVisible = (el) => {
                    if (!el || !el.getBoundingClientRect) return false;
                    const r = el.getBoundingClientRect();
                    const s = window.getComputedStyle(el);
                    return r.width > 0 && r.height > 0 &&
                        s.display !== 'none' && s.visibility !== 'hidden' &&
                        Number(s.opacity || '1') > 0;
                };
                const textOf = (el) => [
                    el.innerText, el.textContent,
                    el.getAttribute('aria-label'),
                    el.getAttribute('title'),
                    el.getAttribute('value')
                ].filter(Boolean).join(' ').replace(/\\s+/g, ' ').trim();
                const clickEl = (el) => {
                    el.scrollIntoView({block: 'center', inline: 'nearest', behavior: 'instant'});
                    if (typeof el.click === 'function') el.click();
                    else el.dispatchEvent(new MouseEvent('click', {bubbles: true, cancelable: true, composed: true, view: window}));
                };
                const start = Array.from(document.querySelectorAll('button,input[type=button],input[type=submit]'))
                    .filter(isVisible)
                    .find(el => /\\bstart\\b/i.test(textOf(el)));
                if (!start) return false;
                clickEl(start);
                return true;
            }"""
        )
    except Exception as exc:
        logger.warning("[RPA CHALLENGE] failed to click Start gate: %s", exc)
        return state
    if clicked:
        browser.rpa_trail.append(
            {
                "action": "click_text",
                "type_value": "Start",
                "method": "rpa_challenge_start_gate",
            }
        )
        logger.info("[RPA CHALLENGE] Start gate clicked; waiting for Round 1")
        await asyncio.sleep(0.8)
        state = await _get_rpa_challenge_state(page)
    return state


async def _run_rpa_challenge_if_present(browser: "BrowserEnv", goal: str) -> bool:
    page = await browser._ensure_active_page(reason="inspect rpachallenge deterministic entry")
    if not page:
        return False
    state = await _get_rpa_challenge_state(page)
    if not state.get("is_challenge"):
        return False
    if _rpa_challenge_is_complete(state):
        logger.info("[RPA CHALLENGE] already in completed state")
        return True
    state = await _start_rpa_challenge_if_needed(browser, page)
    if not state.get("current_round"):
        logger.info(
            "[RPA CHALLENGE] detected but no active round yet; state=%s",
            {k: state.get(k) for k in ("current_round", "start_visible", "submit_visible", "success")},
        )
        return False
    return await _run_rpa_challenge_macro(browser, page, state)


async def _get_round_form_state(page) -> dict:
    """Generic multi-round form detector; legacy RPA Challenge helpers use it too."""
    return await _get_rpa_challenge_state(page)


def _round_form_is_complete(state: dict | None) -> bool:
    return _rpa_challenge_is_complete(state)


async def _has_active_round_form_round(browser: "BrowserEnv") -> bool:
    return await _has_active_rpa_challenge_round(browser)


async def _run_round_form_if_present(browser: "BrowserEnv", goal: str) -> bool:
    return await _run_rpa_challenge_if_present(browser, goal)


async def _run_rpa_challenge_macro(
    browser: "BrowserEnv",
    page,
    state: dict,
) -> bool:
    if not state.get("is_challenge") or not state.get("current_round"):
        return False

    total_rounds = int(state.get("total_rounds") or _RPA_CHALLENGE_TOTAL_ROUNDS)
    download_url = str(state.get("download_url") or "")
    if not download_url and re.search(r"rpachallenge\.com", str(state.get("page_url") or ""), re.I):
        download_url = _RPA_CHALLENGE_DEFAULT_XLSX
    if not download_url:
        logger.info("[ROUND FORM] detected round form but no spreadsheet/csv download link found; leaving to generic form/VLM path")
        return False
    try:
        rows = _load_rpa_challenge_rows(
            download_url,
            total_rounds,
        )
    except Exception as exc:
        logger.warning("[RPA CHALLENGE] unable to load spreadsheet: %s", exc)
        return False
    if len(rows) < total_rounds:
        logger.warning(
            "[RPA CHALLENGE] spreadsheet rows insufficient: %s/%s",
            len(rows),
            total_rounds,
        )
        return False

    logger.info(
        "[RPA CHALLENGE] deterministic macro start from round %s/%s",
        state.get("current_round"),
        total_rounds,
    )
    for round_no in range(int(state.get("current_round") or 1), total_rounds + 1):
        row = rows[round_no - 1]
        result = await page.evaluate(
            """async ({row, expectedRound}) => {
                const isVisible = (el) => {
                    if (!el || !el.getBoundingClientRect) return false;
                    const r = el.getBoundingClientRect();
                    const s = window.getComputedStyle(el);
                    return r.width > 0 && r.height > 0 &&
                        s.display !== 'none' && s.visibility !== 'hidden' &&
                        Number(s.opacity || '1') > 0;
                };
                const clean = (s) => String(s || '').replace(/\\s+/g, ' ').trim();
                const cleanLabel = (s) => clean(s).replace(/^[*\\s:：-]+|[*\\s:：-]+$/g, '');
                const labelTextOf = (el) => cleanLabel(
                    el?.innerText || el?.textContent || el?.getAttribute?.('aria-label') || ''
                );
                const norm = (s) => cleanLabel(s).toLowerCase();
                const textOf = (el) => [
                    el.innerText, el.textContent,
                    el.getAttribute('aria-label'),
                    el.getAttribute('placeholder'),
                    el.getAttribute('title'),
                    el.getAttribute('value'),
                    el.name, el.id
                ].filter(Boolean).join(' ').replace(/\\s+/g, ' ').trim();
                const setNativeValue = (el, value) => {
                    const proto = el instanceof HTMLTextAreaElement ? HTMLTextAreaElement.prototype : HTMLInputElement.prototype;
                    const setter = Object.getOwnPropertyDescriptor(proto, 'value')?.set;
                    if (setter) setter.call(el, value); else el.value = value;
                    el.dispatchEvent(new Event('input', {bubbles: true}));
                    el.dispatchEvent(new Event('change', {bubbles: true}));
                    el.dispatchEvent(new Event('blur', {bubbles: true}));
                };
                const clickEl = (el) => {
                    el.scrollIntoView({block: 'center', inline: 'nearest', behavior: 'instant'});
                    if (typeof el.click === 'function') el.click();
                    else el.dispatchEvent(new MouseEvent('click', {bubbles: true, cancelable: true, composed: true, view: window}));
                };
                const currentRoundMatch = clean(document.body?.innerText || '').match(/ROUND\\s*(\\d+)/i) ||
                    clean(document.body?.innerText || '').match(/第\\s*(\\d+)\\s*(?:轮|回合)/i);
                const currentRound = currentRoundMatch ? Number(currentRoundMatch[1]) : null;
                if (!currentRound || currentRound !== Number(expectedRound)) {
                    return {ok: false, reason: 'round_mismatch', currentRound, expectedRound};
                }

                const controls = Array.from(document.querySelectorAll('input,textarea,select'))
                    .filter(isVisible)
                    .filter(el => !['hidden', 'button', 'submit', 'reset', 'checkbox', 'radio'].includes((el.type || '').toLowerCase()));
                const labels = Array.from(document.querySelectorAll('label,.el-form-item__label,.ant-form-item-label,.n-form-item-label,span,div'))
                    .filter(isVisible);

                const resolveField = (label, used) => {
                    const labelNorm = norm(label);
                    const labelHits = [];
                    for (const el of labels) {
                        const text = labelTextOf(el);
                        if (!text) continue;
                        const textNorm = norm(text);
                        if (textNorm !== labelNorm) continue;
                        const r = el.getBoundingClientRect();
                        let score = 200;
                        if (el.tagName === 'LABEL') score += 300;
                        if (r.width <= 180 && r.height <= 30) score += 40;
                        labelHits.push({el, score});
                    }
                    labelHits.sort((a, b) => b.score - a.score);
                    for (const hit of labelHits.slice(0, 5)) {
                        let cur = hit.el;
                        for (let depth = 0; cur && depth < 6; depth++) {
                            const candidates = controls.filter(ctrl => !used.has(ctrl));
                            const localControls = candidates.filter(ctrl => cur.contains(ctrl));
                            if (localControls.length === 1) {
                                return {labelEl: hit.el, control: localControls[0]};
                            }
                            if (localControls.length > 1) {
                                const lr = hit.el.getBoundingClientRect();
                                const scored = localControls.map(ctrl => {
                                    const cr = ctrl.getBoundingClientRect();
                                    const labelMidY = lr.top + lr.height / 2;
                                    const controlMidY = cr.top + cr.height / 2;
                                    const sameRowGap = Math.abs(controlMidY - labelMidY);
                                    const rightGap = cr.left - lr.right;
                                    const verticalGap = cr.top - lr.bottom;
                                    const horizontalDelta = Math.abs((cr.left + cr.width / 2) - (lr.left + lr.width / 2));
                                    let score = 200 - horizontalDelta - Math.abs(verticalGap) * 2;
                                    if (sameRowGap <= Math.max(32, Math.max(lr.height, cr.height)) && rightGap >= -24) {
                                        score += 260 - sameRowGap - Math.max(0, rightGap) / 20;
                                    }
                                    if (verticalGap >= -12 && verticalGap <= 120) score += 180;
                                    if (Math.abs(cr.left - lr.left) <= 40) score += 60;
                                    return {ctrl, score};
                                }).sort((a, b) => b.score - a.score);
                                if (scored[0]) return {labelEl: hit.el, control: scored[0].ctrl};
                            }
                            cur = cur.parentElement;
                        }
                    }

                    const fallbackLabel = labelHits[0]?.el;
                    if (!fallbackLabel) return null;
                    const lr = fallbackLabel.getBoundingClientRect();
                    const scored = controls
                        .filter(ctrl => !used.has(ctrl))
                        .map(ctrl => {
                            const cr = ctrl.getBoundingClientRect();
                            const labelMidY = lr.top + lr.height / 2;
                            const controlMidY = cr.top + cr.height / 2;
                            const sameRowGap = Math.abs(controlMidY - labelMidY);
                            const rightGap = cr.left - lr.right;
                            const verticalGap = cr.top - lr.bottom;
                            const horizontalDelta = Math.abs((cr.left + cr.width / 2) - (lr.left + lr.width / 2));
                            let score = 200 - horizontalDelta - Math.abs(verticalGap) * 2;
                            if (sameRowGap <= Math.max(32, Math.max(lr.height, cr.height)) && rightGap >= -24) {
                                score += 260 - sameRowGap - Math.max(0, rightGap) / 20;
                            }
                            if (verticalGap >= -12 && verticalGap <= 120) score += 180;
                            if (Math.abs(cr.left - lr.left) <= 40) score += 60;
                            return {ctrl, score};
                        })
                        .sort((a, b) => b.score - a.score);
                    if (!scored[0]) return null;
                    return {labelEl: fallbackLabel, control: scored[0].ctrl};
                };

                const used = new Set();
                const filled = [];
                for (const [label, value] of Object.entries(row || {})) {
                    const binding = resolveField(label, used);
                    if (!binding || !binding.control) {
                        return {ok: false, reason: 'field_not_found', label, filled};
                    }
                    used.add(binding.control);
                    binding.control.scrollIntoView({block: 'center', inline: 'nearest', behavior: 'instant'});
                    setNativeValue(binding.control, String(value || ''));
                    const observed = clean(binding.control.value || binding.control.getAttribute('value') || '');
                    filled.push({label, expected: String(value || ''), observed});
                    if (norm(observed) !== norm(value)) {
                        return {ok: false, reason: 'value_mismatch', label, expected: value, observed, filled};
                    }
                }

                const submit = Array.from(document.querySelectorAll('button,input[type=submit],input[type=button]'))
                    .filter(isVisible)
                    .find(el => /submit|create|save|send|apply|提交|保存|确定/i.test(clean(textOf(el))));
                if (!submit) {
                    return {ok: false, reason: 'submit_not_found', filled};
                }
                clickEl(submit);
                return {
                    ok: true,
                    round: currentRound,
                    submitText: clean(textOf(submit)),
                    filled,
                };
            }""",
            {"row": row, "expectedRound": round_no},
        )
        if not result or not result.get("ok"):
            logger.warning("[RPA CHALLENGE] round %s failed: %s", round_no, result)
            return False
        browser.rpa_trail.append(
            {
                "action": "rpa_challenge_round",
                "round": round_no,
                "total_rounds": total_rounds,
                "method": "spreadsheet_bound_geometry",
                "fields": result.get("filled") or [],
                "submit_text": result.get("submitText") or "Submit",
            }
        )

        await asyncio.sleep(0.8)
        next_state = await _get_rpa_challenge_state(page)
        if round_no < total_rounds:
            if int(next_state.get("current_round") or 0) != round_no + 1:
                logger.warning(
                    "[RPA CHALLENGE] round did not advance after submit: expected=%s got=%s",
                    round_no + 1,
                    next_state.get("current_round"),
                )
                return False
        else:
            if not _rpa_challenge_is_complete(next_state):
                logger.warning("[RPA CHALLENGE] final state ambiguous after round %s: %s", round_no, next_state)
                return False

    print("\033[1;32m✅ [RPA CHALLENGE]\033[0m Completed deterministic rounds with spreadsheet-driven mapping")
    _broadcast_log_safe("[RPA CHALLENGE] Deterministic spreadsheet-driven rounds completed")
    return True

